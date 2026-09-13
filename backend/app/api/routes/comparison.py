import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlmodel import Session

from ...core.db import get_session
from ...core.security import verify_api_key
from ...schemas.comparison import CompareRequest, CompareResponse
from ...services.comparison import compare_documents

router = APIRouter(prefix="/documents", tags=["comparison"])


@router.post("/compare", response_model=CompareResponse)
async def compare_documents_endpoint(
    body: CompareRequest,
    session: Session = Depends(get_session),
    _: str = Depends(verify_api_key),
) -> CompareResponse:
    try:
        return compare_documents(body.document_a, body.document_b, session)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get("/compare/export.xlsx")
async def compare_export_excel(
    document_a: str,
    document_b: str,
    session: Session = Depends(get_session),
    _: str = Depends(verify_api_key),
) -> StreamingResponse:
    try:
        result = compare_documents(document_a, document_b, session)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Count"])
    for k, v in result.summary.model_dump().items():
        ws.append([k.replace("_", " ").title(), v])

    ws2 = wb.create_sheet("Chunk Changes")
    ws2.append(["Change Type", "Doc A Text", "Doc B Text"])
    for c in result.chunk_changes:
        ws2.append([c.type, c.text_a or "", c.text_b or ""])

    ws3 = wb.create_sheet("Signal Changes")
    ws3.append(["Signal Type", "Change", "Value A", "Value B", "Evidence A", "Evidence B"])
    for s in result.signal_changes:
        ws3.append([
            s.signal_type, s.change_type,
            s.value_a or s.value or "", s.value_b or "",
            s.evidence, s.evidence_b or "",
        ])

    ws4 = wb.create_sheet("Entity Changes")
    ws4.append(["Entity Type", "Change", "Value A", "Value B", "Evidence A", "Evidence B"])
    for e in result.entity_changes:
        ws4.append([
            e.signal_type, e.change_type,
            e.value_a or e.value or "", e.value_b or "",
            e.evidence, e.evidence_b or "",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=comparison.xlsx"},
    )
